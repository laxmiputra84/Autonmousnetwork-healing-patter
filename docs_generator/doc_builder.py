import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOC_DIR = r"d:\Autonomous Network Healing Platform\docs_generator"
DOCX_PATH = os.path.join(DOC_DIR, "Autonomous_Network_Healing_Platform_Documentation.docx")
PDF_PATH = os.path.join(DOC_DIR, "Autonomous_Network_Healing_Platform_Documentation.pdf")
FIG1 = os.path.join(DOC_DIR, "figure1_architecture.png")
FIG2 = os.path.join(DOC_DIR, "figure2_closed_loop.png")
FIG3 = os.path.join(DOC_DIR, "figure3_er_diagram.png")

# Palette constants
COLOR_NAVY = RGBColor(15, 23, 42)       # #0F172A
COLOR_SLATE = RGBColor(51, 65, 85)      # #334155
COLOR_BLUE = RGBColor(2, 132, 199)      # #0284C7
COLOR_GRAY = RGBColor(100, 116, 139)    # #64748B
HEX_NAVY = "0F172A"
HEX_SLATE_HEADER = "1E293B"
HEX_LIGHT_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_BLUE_BG = "F0F9FF"
HEX_BLUE_BORDER = "0284C7"

def add_footer_field(p, field_name):
    """Safely append a Word field (PAGE, NUMPAGES) using separate runs."""
    fld1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instr = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> {field_name} </w:instrText>')
    fld2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fld3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    p.add_run()._r.append(fld1)
    p.add_run()._r.append(instr)
    p.add_run()._r.append(fld2)
    p.add_run()._r.append(fld3)

def create_document():
    doc = Document()
    
    # Configure A4 page setup and 1.0 inch margins
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)
        
        # Configure Header & Footer
        section.different_first_page_header_footer = True
        
        # Header (Pages 2-12)
        hdr = section.header
        p_hdr = hdr.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.paragraph_format.space_before = Pt(0)
        p_hdr.paragraph_format.space_after = Pt(0)
        r_hdr = p_hdr.add_run("Autonomous Network Healing Platform  |  Technical Documentation")
        r_hdr.font.name = 'Arial'
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = COLOR_GRAY
        
        # Footer (Pages 2-12)
        ftr = section.footer
        p_ftr = ftr.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_ftr.paragraph_format.space_before = Pt(0)
        p_ftr.paragraph_format.space_after = Pt(0)
        
        r_ftr_l = p_ftr.add_run("Confidential  •  Academic & Technical Documentation Report\t\tPage ")
        r_ftr_l.font.name = 'Arial'
        r_ftr_l.font.size = Pt(8.5)
        r_ftr_l.font.color.rgb = COLOR_GRAY
        
        # Dynamic Page Field
        add_footer_field(p_ftr, "PAGE")
        
        r_ftr_mid = p_ftr.add_run(" of ")
        r_ftr_mid.font.name = 'Arial'
        r_ftr_mid.font.size = Pt(8.5)
        r_ftr_mid.font.color.rgb = COLOR_GRAY
        
        # Dynamic NumPages Field
        add_footer_field(p_ftr, "NUMPAGES")
        
    return doc

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def add_h1(doc, text, space_before=6, space_after=2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = COLOR_NAVY
    return p

def add_h2(doc, text, space_before=4, space_after=1.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_BLUE
    return p

def add_h3(doc, text, space_before=3, space_after=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = COLOR_SLATE
    return p

def add_p(doc, text, bold_prefix="", space_after=1.5, line_spacing=1.12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Arial'
        r_b.font.size = Pt(9.5)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_NAVY
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_SLATE
    return p

def add_bullet(doc, bold_prefix, text, space_after=1.2, line_spacing=1.12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    r_dot = p.add_run("• ")
    r_dot.font.name = 'Arial'
    r_dot.font.size = Pt(9.5)
    r_dot.font.bold = True
    r_dot.font.color.rgb = COLOR_BLUE
    if bold_prefix:
        r_b = p.add_run(bold_prefix + " ")
        r_b.font.name = 'Arial'
        r_b.font.size = Pt(9.5)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_NAVY
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_SLATE
    return p

def add_callout(doc, title, text, bg_hex="F0F9FF", border_hex="0284C7", **kwargs):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.27)
    set_cell_margins(cell, top=50, bottom=50, left=100, right=100)
    set_cell_shading(cell, bg_hex)
    set_cell_border(cell, 
                    left=dict(val='single', sz=20, color=border_hex),
                    top=dict(val='none'), right=dict(val='none'), bottom=dict(val='none'))
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r_t = p.add_run(f"■ {title}\n")
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(9.0)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_NAVY
    
    r_b = p.add_run(text)
    r_b.font.name = 'Arial'
    r_b.font.size = Pt(8.5)
    r_b.font.color.rgb = COLOR_SLATE
    return tbl

def add_table_data(doc, col_widths, headers, data, font_size=7.5, **kwargs):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    hdr_row = tbl.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for i, title in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = Inches(col_widths[i])
        set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
        set_cell_shading(cell, HEX_SLATE_HEADER)
        set_cell_border(cell, 
                        top=dict(val='single', sz=6, color='0F172A'),
                        bottom=dict(val='single', sz=12, color='0F172A'),
                        left=dict(val='none'), right=dict(val='none'))
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.font.name = 'Arial'
        r.font.size = Pt(font_size)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for row_idx, row_data in enumerate(data):
        row = tbl.rows[row_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg_col = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        for col_idx, val in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.width = Inches(col_widths[col_idx])
            set_cell_margins(cell, top=35, bottom=35, left=80, right=80)
            set_cell_shading(cell, bg_col)
            set_cell_border(cell, 
                            bottom=dict(val='single', sz=4, color='E2E8F0'),
                            top=dict(val='none'), left=dict(val='none'), right=dict(val='none'))
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            r = p.add_run(str(val))
            r.font.name = 'Arial'
            r.font.size = Pt(font_size - 0.3)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = COLOR_NAVY
            else:
                r.font.color.rgb = COLOR_SLATE
    return tbl

def add_figure(doc, img_path, caption, width_in=5.8, **kwargs):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    r.add_picture(img_path, width=Inches(width_in))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(2)
    p_cap.paragraph_format.keep_with_next = False
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = 'Arial'
    r_cap.font.size = Pt(8.0)
    r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_GRAY
