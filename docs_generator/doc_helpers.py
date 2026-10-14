import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOC_DIR = r"d:\Autonomous Network Healing Platform\docs_generator"
DOCX_PATH = os.path.join(DOC_DIR, "Autonomous_Network_Healing_Platform_Documentation.docx")
PDF_PATH = os.path.join(DOC_DIR, "Autonomous_Network_Healing_Platform_Documentation.pdf")
FIG1 = os.path.join(DOC_DIR, "figure1_architecture.png")
FIG2 = os.path.join(DOC_DIR, "figure2_closed_loop.png")
FIG3 = os.path.join(DOC_DIR, "figure3_er_diagram.png")

# Palette constants
COLOR_NAVY = RGBColor(15, 23, 42)       # #0F172A
COLOR_SLATE = RGBColor(51, 65, 85)      # #334155
COLOR_BLUE = RGBColor(2, 132, 199)      # #0284C7
COLOR_GRAY = RGBColor(100, 116, 139)    # #64748B
HEX_PRIMARY = "0F172A"
HEX_LIGHT_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_BLUE_BG = "F0F9FF"
HEX_BLUE_BORDER = "0284C7"

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell margins in dxa (1 pt = 20 dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='0284C7')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def add_callout(doc, title, text, bg_hex="F0F9FF", border_hex="0284C7"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.27)
    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
    set_cell_shading(cell, bg_hex)
    set_cell_border(cell, 
                    left=dict(val='single', sz=24, color=border_hex),
                    top=dict(val='none'), right=dict(val='none'), bottom=dict(val='none'))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"■ {title}\n")
    run_t.font.name = 'Arial'
    run_t.font.size = Pt(9.5)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_NAVY
    
    run_b = p.add_run(text)
    run_b.font.name = 'Arial'
    run_b.font.size = Pt(9)
    run_b.font.color.rgb = COLOR_SLATE
    
    # Spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(2)
    return tbl

def format_table(tbl, col_widths, headers, data, font_size=8.5):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Header Row
    hdr_row = tbl.rows[0]
    # CantSplit and tblHeader
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for i, title in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = Inches(col_widths[i])
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        set_cell_shading(cell, "1E293B")
        set_cell_border(cell, 
                        top=dict(val='single', sz=6, color='0F172A'),
                        bottom=dict(val='single', sz=12, color='0F172A'),
                        left=dict(val='none'), right=dict(val='none'))
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.font.name = 'Arial'
        r.font.size = Pt(font_size)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    
    # Data Rows
    for row_idx, row_data in enumerate(data):
        row = tbl.rows[row_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg_col = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        for col_idx, val in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.width = Inches(col_widths[col_idx])
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            set_cell_shading(cell, bg_col)
            set_cell_border(cell, 
                            bottom=dict(val='single', sz=4, color='E2E8F0'),
                            top=dict(val='none'), left=dict(val='none'), right=dict(val='none'))
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(val))
            r.font.name = 'Arial'
            r.font.size = Pt(font_size - 0.5)
            r.font.color.rgb = COLOR_NAVY if col_idx == 0 else COLOR_SLATE

print("Helper functions defined successfully.")
