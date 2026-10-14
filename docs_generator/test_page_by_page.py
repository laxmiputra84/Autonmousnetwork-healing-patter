import docx, win32com.client, os, sys
import doc_builder

word = win32com.client.Dispatch("Word.Application")
word.Visible = False
word.DisplayAlerts = 0

def try_open(doc, name):
    p = os.path.abspath(name)
    doc.save(p)
    try:
        d = word.Documents.Open(p)
        pg = d.ComputeStatistics(2)
        d.Close(False)
        print(f"Page test {name}: PASSED, pages={pg}")
        return True
    except Exception as e:
        print(f"Page test {name}: FAILED, error={e}")
        return False

try:
    d = doc_builder.create_document()
    p_title = d.add_paragraph()
    p_title.add_run("Title")
    try_open(d, "p1.docx")
    
    d.add_page_break()
    p2 = d.add_paragraph()
    p2.add_run("Page 2 content")
    try_open(d, "p2.docx")
    
    # Test table
    headers = ["Col 1", "Col 2", "Col 3"]
    data = [["A", "B", "C"], ["D", "E", "F"]]
    doc_builder.add_table_data(d, [2.0, 2.0, 2.27], headers, data)
    try_open(d, "p2_tbl.docx")
    
    # Test Figure 1
    d.add_page_break()
    doc_builder.add_figure(d, doc_builder.FIG1, "Fig 1", width_in=5.8)
    try_open(d, "p_fig1.docx")
    
    # Test Figure 2
    d.add_page_break()
    doc_builder.add_figure(d, doc_builder.FIG2, "Fig 2", width_in=5.8)
    try_open(d, "p_fig2.docx")
    
    # Test Figure 3
    d.add_page_break()
    doc_builder.add_figure(d, doc_builder.FIG3, "Fig 3", width_in=5.8)
    try_open(d, "p_fig3.docx")
    
finally:
    word.Quit()
