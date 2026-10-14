import docx, win32com.client, os
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def add_footer_field(p, field_name):
    fld1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instr = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> {field_name} </w:instrText>')
    fld2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fld3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    p.add_run()._r.append(fld1)
    p.add_run()._r.append(instr)
    p.add_run()._r.append(fld2)
    p.add_run()._r.append(fld3)

doc = docx.Document()
sec = doc.sections[0]
sec.different_first_page_header_footer = True
p = sec.footer.paragraphs[0]
p.add_run('Confidential | Page ')
add_footer_field(p, 'PAGE')
p.add_run(' of ')
add_footer_field(p, 'NUMPAGES')

doc.save('test_field.docx')

word = win32com.client.Dispatch('Word.Application')
word.Visible = False
word.DisplayAlerts = 0
try:
    d = word.Documents.Open(os.path.abspath('test_field.docx'))
    print('Test field pages:', d.ComputeStatistics(2))
    d.Close(False)
    print('Test field SUCCESS!')
finally:
    word.Quit()
