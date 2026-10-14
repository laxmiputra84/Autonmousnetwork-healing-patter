import docx, win32com.client, os

def test_doc(doc, name):
    path = os.path.abspath(name)
    doc.save(path)
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(path)
        p = d.ComputeStatistics(2)
        d.Close(False)
        print(f"{name}: SUCCESS, pages={p}")
    except Exception as e:
        print(f"{name}: FAILED, error={e}")
    finally:
        word.Quit()

# Test 1: Just paragraphs
doc1 = docx.Document()
doc1.add_paragraph("Test 1")
test_doc(doc1, "test1.docx")

# Test 2: Table with custom XML shading/borders
from doc_builder import set_cell_margins, set_cell_shading, set_cell_border
doc2 = docx.Document()
tbl = doc2.add_table(rows=2, cols=2)
set_cell_shading(tbl.cell(0,0), "0F172A")
set_cell_margins(tbl.cell(0,0))
set_cell_border(tbl.cell(0,0), bottom=dict(val='single', sz=4, color='CBD5E1'))
tbl.cell(0,0).paragraphs[0].add_run("Table test")
test_doc(doc2, "test2.docx")

# Test 3: Picture
doc3 = docx.Document()
doc3.add_picture(r"d:\Autonomous Network Healing Platform\docs_generator\figure1_architecture.png", width=docx.shared.Inches(5))
test_doc(doc3, "test3.docx")

# Test 4: Header/footer without dynamic fields
doc4 = docx.Document()
sec = doc4.sections[0]
sec.different_first_page_header_footer = True
sec.header.paragraphs[0].add_run("Header test")
sec.footer.paragraphs[0].add_run("Footer test Page 1 of 12")
test_doc(doc4, "test4.docx")
